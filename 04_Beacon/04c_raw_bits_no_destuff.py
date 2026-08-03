# -*- coding: utf-8 -*-
"""
04c_raw_bits_no_destuff.py -- exactly `04_beacon_field_decode.py`'s pipeline
(same theta*/yc decision rule, same red "near decision line" flagging, same
field table / header rows / FCS row / trailing-flag rows / alignment check),
with ONE change: the destuffing step is skipped entirely. Bits that would
normally get removed as "stuffed 0s" are left in the stream.

Why this exists: 04_beacon_field_decode.py's field table depends on
destuffing to place byte boundaries correctly. The frame-alignment check
already proved that destuffing doesn't cut all 3 cut_first3.ogg frames
exactly at the closing flag (see 04_Beacon/README.md "Frame-alignment
self-check") -- but that check still trusts the destuffing code's own
bookkeeping (how many bits it removed, and where). This script is the
independent cross-check: skip destuffing altogether, and see what "byte N"
looks like using nothing but the raw threshold=0/theta* slicer decisions.
Early bytes (before any real HDLC stuffing has occurred) should still line
up with the destuffed version and with REFERENCE_FRAME.md; once real
stuffing starts, this raw view increasingly diverges from both -- and by
how much, measured directly, is the answer to "how many bits offset, if we
don't trust destuffing at all".

This is intentionally a copy of 04_beacon_field_decode.py with the minimum
edit needed (destuff_with_map's call site removed, decode_frame just passes
the raw arrays straight through), not a from-scratch reimplementation --
see CLAUDE.md's self-contained-script convention. Every decision rule,
color, and column below is identical to that file; only search for
"NO DESTUFF" comments to see what actually changed.

Usage: `python 04c_raw_bits_no_destuff.py [audio.ogg] [--frame N] [--z-threshold Z]`.
Same CLI as 04_beacon_field_decode.py.

Output: `Output/<audio stem>_frame<N>_raw_no_destuff.xlsx` -- a DIFFERENT
filename from `04_beacon_field_decode.py`'s own output, so running this never
overwrites the regular decode.
"""

import argparse
import os
import re
import sys
import json

import numpy as np
import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Font, PatternFill, Alignment

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
sys.path.insert(0, HERE)

from scionx import audio_io, baseline               # noqa: E402
from scionx.hdlc import crc16_x25                    # noqa: E402

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:  # noqa: BLE001
    pass

FS, SPS = 48000, 5
PHASE = 2
AUDIO = os.path.join(os.path.dirname(HERE), "Data", "cut_first3.ogg")
XLSX_PATH = os.path.join(HERE, "SCIONX_TLMnew.xlsx")
ENUMS_PATH = os.path.join(HERE, "SCIONX_enums.json")
OUT_DIR = os.path.join(HERE, "Output")

FLAGS4 = bytes([0x7E] * 4)
HEADER_ADDR = bytes.fromhex("849c6086aa4060849c60a686b0e1")   # Dest+Src address, 14 bytes

Z_THRESHOLD = 12.0
MIN_FRAME_GAP = 100_000
SEARCH_SAMPLES = 19211    # header through the trailing flags, same window 01/02/03 use
MARGIN_FRAC = 0.10        # "near decision line" half-band, fraction of A=median(|y|) -- same as 03

INFO_BYTE0 = 16           # Info field starts at payload byte 16 (after Dest+Src+Control+PID)
INFO_NBYTES = 256         # REFERENCE_FRAME.md: Info = 0x10-0x10F = 256 bytes = 2048 bits (matches xlsx's max OffsetBit+BitLen)

REF_ROWS = (
    "849c6086aa4060849c60a686b0e103f0"
    "007c083c81336a000201000000000020"
    "0000002b81336a000000000000030200"
    "0044090020100000c3add00099bf4e04"
    "8bc5780561d7f608d4a11f0000000000"
    "00000000000000000000000000000000"
    "00000000000000000000000000000000"
    "00000000000000000000000000000000"
    "00000000000000000000000000000000"
    "00000000000000000000000000000826"
    "08260826282620262000200020002000"
    "2000a03d78ff770f9bff8a1f161515e7"
    "b0b0b0b019181ae80c4002481b180000"
    "00000000000000000000000000000000"
    "00000000000000000000000000000000"
    "00000000000000000000000000000000"
    "00000000000000000000000000000000"
)
REF = bytes.fromhex(REF_ROWS)


# ============================================================================
# Signal-side helpers, duplicated from 04_beacon_field_decode.py (this repo's
# convention: every script is self-contained, see CLAUDE.md)
# ============================================================================

def bits_lsb_first(data):
    out = []
    for b in data:
        for k in range(8):
            out.append((b >> k) & 1)
    return out


def build_header_template():
    bits = np.array(bits_lsb_first(FLAGS4) + bits_lsb_first(HEADER_ADDR), dtype=np.float64)
    nrz = 2 * bits - 1
    return np.repeat(nrz, SPS)


def detect_frame_starts(yc, template):
    L = template.size
    corr = np.correlate(yc, template, mode="valid")
    csum2 = np.concatenate(([0.0], np.cumsum(yc.astype(np.float64) ** 2)))
    window_energy = csum2[L:] - csum2[:-L]
    template_norm = np.sqrt(np.sum(template ** 2))
    R = corr / (template_norm * np.sqrt(window_energy) + 1e-12)

    z = R / np.std(R)
    candidates = np.where(z > Z_THRESHOLD)[0]
    candidates = candidates[np.argsort(-z[candidates])]

    starts = []
    for idx in candidates:
        if all(abs(int(idx) - s) > MIN_FRAME_GAP for s in starts):
            starts.append(int(idx))
    order = np.argsort(starts)
    starts = [starts[i] for i in order]
    zs = [float(z[s]) for s in starts]
    return starts, zs


def find_zero_runs(buf, min_len=8):
    runs, i = [], 0
    while i < len(buf):
        if buf[i] == 0:
            j = i
            while j < len(buf) and buf[j] == 0:
                j += 1
            if j - i >= min_len:
                runs.append((i, j))
            i = j
        else:
            i += 1
    return runs


def bits_to_bytes(bits):
    n = bits.size // 8
    out = bytearray()
    for i in range(n):
        v = 0
        for k in range(8):
            v |= int(bits[i * 8 + k]) << k
        out.append(v)
    return bytes(out)


def build_segment_map():
    """Byte-range map (address/data1/zero-run1/data2/zero-run2/FCS), derived
    from REFERENCE_FRAME.md's known zero-runs -- identical logic to
    04_beacon_field_decode.py. NO DESTUFF: applied here directly to raw,
    un-destuffed bit positions, so "byte N" only means "the true payload byte
    N" for as long as no real stuffing has occurred before it yet."""
    zero_runs = find_zero_runs(REF)
    segments = [("address", 0, 14)]
    prev = 14
    seg_id = 1
    for (z0, z1) in zero_runs:
        if z0 > prev:
            segments.append((f"data{seg_id}", prev, z0))
        segments.append((f"zero-run{seg_id}", z0, z1))
        prev = z1
        seg_id += 1
    if prev < 272:
        segments.append((f"data{seg_id}", prev, 272))
    segments.append(("FCS", 272, 274))
    return segments


FLAG_BITS = [0, 1, 1, 1, 1, 1, 1, 0]     # 0x7E as it appears on the wire
FRAME_END_BIT = 274 * 8                  # 272-byte payload + 2-byte FCS, IF there were zero stuffed bits
TRAIL_BYTES = 16                         # bytes of trailing-flag region to decode and show


def find_flag_run(bits, lo, hi, min_run=3):
    """First position p in [lo,hi) that starts a run of `min_run` back-to-back
    0x7E flags. Requiring a *run* (not a single match) matters: an isolated
    01111110 can occur by chance inside the payload, but flags spaced exactly
    8 bits apart cannot.

    KNOWN LIMITATION (found by user inspection, cut_first3.ogg frame#1/#3):
    this exact-match requirement is brittle against a single low-confidence
    bit. If the very first real flag copy has one bit mis-sliced (near the
    decision line), it fails the exact-match test, so the scan skips past it
    and locks onto a LATER run instead -- overshooting the true offset by a
    full flag period (8 bits) or more. Concretely: frame#1's real flag starts
    at bit 2196 with 1 bit off (ndiff=1 vs the clean 01111110), but this
    function returns 2204 (the next fully-clean run) -- an offset of +12
    instead of the correct +4. Frame#3 similarly returns +21 instead of the
    correct +5 (true start: a single clean match at bit 2197, immediately
    followed by another 1-bit-off copy that breaks the 3-in-a-row requirement).

    Not fixed here on purpose: telling "a real flag with one wrong bit" apart
    from "not a flag at all" is exactly the kind of judgment call that needs
    a human looking at context (the write_workbook() trailing-flag rows, with
    each byte's diff-from-0x7E and near-decision-line marks, are meant for
    this). Treat this function's return value as a rough/conservative
    estimate -- always eyeball the trailing rows before trusting an exact
    offset count."""
    seq = bits.tolist() if hasattr(bits, "tolist") else list(bits)
    n = len(seq)
    for p in range(max(0, lo), min(hi, n - 8 * min_run + 1)):
        if all(seq[p + 8 * k: p + 8 * k + 8] == FLAG_BITS for k in range(min_run)):
            return p
    return None


def trailing_expected_pattern(align):
    """The 8-bit pattern a trailing flag byte should show at raw bit
    FRAME_END_BIT, given the detected offset. Same idea as
    04_beacon_field_decode.py's version, just renamed since there's no
    "slip" here (nothing was destuffed to slip) -- just a raw phase offset."""
    off = align.get("offset_bits")
    phase = (-off) % 8 if off is not None else 0
    rot = FLAG_BITS[phase:] + FLAG_BITS[:phase]
    return "".join(map(str, rot))


def check_raw_offset(raw_bits, segments):
    """NO DESTUFF version of 04_beacon_field_decode.py's check_frame_alignment:
    the same structural anchor (closing flag must start at FRAME_END_BIT), but
    with nothing removed from the stream, so `offset_bits` is the TOTAL number
    of bits actually stuffed onto the wire before the closing flag -- measured
    directly from the raw signal, independent of any destuffing code at all.
    There is no "spurious removal" concept here (nothing was removed), so
    unlike the destuffed version this can't localize *where* the real
    stuffing happened -- only how much of it there was in total."""
    n_avail = raw_bits.size
    flag_run_start, offset_bits = None, None
    if n_avail >= FRAME_END_BIT + 8:
        flag_run_start = find_flag_run(raw_bits, FRAME_END_BIT - 48, FRAME_END_BIT + 256)
        if flag_run_start is not None:
            offset_bits = flag_run_start - FRAME_END_BIT
    return {
        "flag_run_start_bit": flag_run_start,
        "offset_bits": offset_bits,
        "aligned": offset_bits == 0,
    }


def calibrate_offset_threshold(y, start):
    t_bits = np.array(bits_lsb_first(FLAGS4) + bits_lsb_first(HEADER_ADDR), dtype=np.uint8)
    idx_hdr = start + SPS * np.arange(0, 144) + PHASE
    y_hdr = y[idx_hdr]
    v1, v0 = y_hdr[t_bits == 1], y_hdr[t_bits == 0]
    grid = np.linspace(min(v0.min(), v1.min()), max(v0.max(), v1.max()), 2000)
    errs = [(v1 <= t).sum() + (v0 > t).sum() for t in grid]
    thr = grid[int(np.argmin(errs))]
    return float(thr), int(min(errs))


def decode_frame(y, yc, start):
    """Identical to 04_beacon_field_decode.py's decode_frame, EXCEPT the
    destuff_with_map() call is skipped -- see the "NO DESTUFF" comment below.
    Every decision rule after that point (theta* for address/data/FCS, yc>0
    for zero-run, the same "near decision line" ambiguous flag) is unchanged."""
    offset_thr, hdr_err = calibrate_offset_threshold(y, start)
    A = float(np.median(np.abs(y)))
    margin = MARGIN_FRAC * A

    span_j1 = int(SEARCH_SAMPLES / SPS)
    idx_all = start + SPS * np.arange(0, span_j1) + PHASE
    idx_all = idx_all[idx_all < y.size]
    body_bits_raw = (y[idx_all[32:]] > 0).astype(np.uint8)     # skip the 4 leading flags; threshold=0 structure pass
    body_idx_raw = idx_all[32:]

    # NO DESTUFF: 04_beacon_field_decode.py calls destuff_with_map() here to
    # remove stuffed 0s and get clean payload-bit indexing. This script skips
    # that entirely -- body_bits_full is just body_bits_raw, unmodified.
    body_bits_full, body_idx_full = body_bits_raw, body_idx_raw

    segments = build_segment_map()
    align = check_raw_offset(body_bits_full, segments)

    n_bits = min(body_bits_full.size, FRAME_END_BIT + TRAIL_BYTES * 8)
    body_bits = body_bits_full[:n_bits]
    body_sample_idx = body_idx_full[:n_bits]

    seg_label = np.empty(n_bits, dtype=object)
    for label, b0, b1 in segments:
        bit0, bit1 = b0 * 8, min(b1 * 8, n_bits)
        if bit0 < bit1:
            seg_label[bit0:bit1] = label
    if n_bits > FRAME_END_BIT:
        seg_label[FRAME_END_BIT:n_bits] = "trailing-flag"

    is_zero_run = np.array([lbl is not None and lbl.startswith("zero-run") for lbl in seg_label])
    is_theta = ~is_zero_run    # address + data1/data2 + FCS: same theta*-decision domain as 04's main script

    y_at_bit = y[body_sample_idx]
    yc_at_bit = yc[body_sample_idx]
    decided_bit = body_bits.copy()
    decided_bit[is_theta] = (y_at_bit[is_theta] > offset_thr).astype(np.uint8)
    decided_bit[is_zero_run] = (yc_at_bit[is_zero_run] > 0).astype(np.uint8)   # 02's Method 1, on yc

    ambiguous = np.zeros(n_bits, dtype=bool)
    ambiguous[is_theta] = np.abs(y_at_bit[is_theta] - offset_thr) <= margin
    ambiguous[is_zero_run] = decided_bit[is_zero_run] == 1

    # informational CRC checks -- essentially guaranteed to fail here even if
    # the destuffed version ever passed, since real stuffed bits are still in
    # the stream; kept anyway for direct comparison against the destuffed sheet.
    crc_baseline = crc_mixed = None
    if n_bits >= 274 * 8:
        fcs_bytes = bits_to_bytes(body_bits[272 * 8:274 * 8])
        crc_baseline = (crc16_x25(bits_to_bytes(body_bits[:272 * 8])).to_bytes(2, "little") == fcs_bytes)
        crc_mixed = (crc16_x25(bits_to_bytes(decided_bit[:272 * 8])).to_bytes(2, "little") == fcs_bytes)

    seg_counts = {}
    for lbl in ("address", "data1", "zero-run1", "data2", "zero-run2", "FCS"):
        mask = seg_label == lbl
        seg_counts[lbl] = (int(ambiguous[mask].sum()), int(mask.sum()))

    return {
        "decided_bit": decided_bit, "ambiguous": ambiguous, "seg_label": seg_label,
        "offset_thr": offset_thr, "hdr_err": hdr_err, "A": A, "margin": margin,
        "crc_baseline": crc_baseline, "crc_mixed": crc_mixed, "n_bits": n_bits,
        "seg_counts": seg_counts, "align": align,
        "bits_thr0": body_bits,
    }


# ============================================================================
# Beacon field layout (xlsx) + enum/transform lookup (json) -- unchanged from
# 04_beacon_field_decode.py
# ============================================================================

def load_beacon_fields(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip() if h else "" for h in rows[0]]
    col = {name: i for i, name in enumerate(header)}

    fields = []
    subsystem = ""
    for r in rows[1:]:
        item_name = r[col["ItemName"]]
        if item_name is None:
            continue
        if r[col["Subsystem"]]:
            subsystem = str(r[col["Subsystem"]]).strip()
        fields.append({
            "subsystem": subsystem,
            "item_name": str(item_name).strip(),
            "dtype": (r[col["DataType"]] or "").strip(),
            "bitlen": int(r[col["BitLen"]]),
            "offsetbit": int(r[col["OffsetBit"]]),
            "endian": (r[col["Endian"]] or "BE").strip().upper(),
            "longdesc": (str(r[col["LongDescription"]]).strip() if r[col["LongDescription"]] else ""),
        })
    fields.sort(key=lambda f: f["offsetbit"])
    return fields


def normalize_name(name):
    return re.sub(r"\s+", "_", name.strip())


def split_top_level(s, sep):
    parts, depth, cur = [], 0, ""
    for ch in s:
        if ch == "(":
            depth += 1
        elif ch == ")":
            depth = max(0, depth - 1)
        if ch == sep and depth == 0:
            parts.append(cur)
            cur = ""
        else:
            cur += ch
    parts.append(cur)
    return parts


def parse_enum_from_longdesc(longdesc):
    if not longdesc:
        return {}
    lines = [ln for ln in longdesc.split("\n") if ln.strip()]
    chunks = []
    if len(lines) > 1:
        chunks = lines
    elif "/" in longdesc and longdesc.count(":") >= 2:
        chunks = longdesc.split("/")
    else:
        chunks = split_top_level(longdesc, ",")

    out = {}
    pat = re.compile(r"^\s*([0-9A-Fa-f]{1,4})\s*[:\-]\s*(.+?)\s*$")
    for chunk in chunks:
        m = pat.match(chunk)
        if not m:
            continue
        key_str, label = m.group(1), m.group(2)
        try:
            key = int(key_str, 16) if re.search(r"[A-Fa-f]", key_str) else int(key_str, 10)
        except ValueError:
            continue
        out[key] = label
    return out


def resolve_lookup(item_name, longdesc, bitlen, enums_ctx):
    norm = normalize_name(item_name)
    norm_lower = norm.lower()

    parsed = parse_enum_from_longdesc(longdesc)
    if parsed:
        return parsed, None, longdesc

    for rule in enums_ctx["nameRules"]:
        if re.match(rule["pattern"], norm):
            if "transform" in rule and rule["transform"] in enums_ctx["transforms"]:
                tr = enums_ctx["transforms"][rule["transform"]]
                return None, tr, f"transform:{rule['transform']} {tr}"
            if "enum" in rule:
                enum_name = rule["enum"]
                table = enums_ctx["enums"].get(enum_name) or enums_ctx["enums"].get(norm)
                if table:
                    table_int = {int(k, 16) if re.search(r"[A-Fa-f]", k) else int(k): v
                                 for k, v in table.items()}
                    return table_int, None, ", ".join(f"{k}={v}" for k, v in table.items())

    for key, table in enums_ctx["enums"].items():
        if key.lower() == norm_lower:
            table_int = {int(k, 16) if re.search(r"[A-Fa-f]", k) else int(k): v
                         for k, v in table.items()}
            return table_int, None, ", ".join(f"{k}={v}" for k, v in table.items())

    for key, tr in enums_ctx["transforms"].items():
        if key.lower() == norm_lower:
            return None, tr, f"transform:{key} {tr}"

    if bitlen == 1:
        default = {0: "OFF", 1: "ON"}
        return default, None, "(assumed 0=OFF/1=ON, no field-specific description)"

    return None, None, "-"


def decode_raw_value(bits_str, dtype, bitlen, endian):
    """Same as 04_beacon_field_decode.py -- deliberately kept identical
    (including its MSB-first bit-order convention) so this sheet's numbers
    are comparable to that one's, with destuffing as the only variable."""
    if bitlen % 8 == 0 and bitlen > 8 and endian == "LE":
        byte_chunks = [bits_str[i:i + 8] for i in range(0, bitlen, 8)]
        bits_str = "".join(reversed(byte_chunks))

    raw_unsigned = int(bits_str, 2)

    if dtype == "float" and bitlen == 32:
        import struct
        as_bytes = raw_unsigned.to_bytes(4, "big")
        return struct.unpack(">f", as_bytes)[0], raw_unsigned

    if dtype.startswith("int") and (raw_unsigned & (1 << (bitlen - 1))):
        return raw_unsigned - (1 << bitlen), raw_unsigned

    return raw_unsigned, raw_unsigned


def format_readable(value, raw_unsigned, parsed_enum, transform):
    if parsed_enum is not None:
        return str(parsed_enum.get(raw_unsigned, f"{raw_unsigned} (unlisted)"))
    if transform is not None:
        scale = transform.get("scale")
        unit = transform.get("unit", "")
        v = value * scale if scale is not None else value
        if isinstance(v, float):
            return f"{v:.4g} {unit}".strip()
        return f"{v} {unit}".strip()
    if isinstance(value, float):
        return f"{value:.6g}"
    return str(value)


# ============================================================================
# Output -- same layout/colors as 04_beacon_field_decode.py
# ============================================================================

RED = InlineFont(color="FFCC0000")
HEADER_FILL = PatternFill("solid", fgColor="FFDCE9F5")

SEGMENT_FILLS = {
    "header":    PatternFill("solid", fgColor="FFDCEAF7"),
    "data1":     PatternFill("solid", fgColor="FFFDEBD3"),
    "zero-run1": PatternFill("solid", fgColor="FFECECEC"),
    "data2":     PatternFill("solid", fgColor="FFFBE0E0"),
    "zero-run2": PatternFill("solid", fgColor="FFE1E1E1"),
    "FCS":       PatternFill("solid", fgColor="FFE6DFF7"),
    "trailing-flag": PatternFill("solid", fgColor="FFDFF3E6"),
    "mixed":     PatternFill("solid", fgColor="FFFFF3B0"),
}


def bits_richtext(bits_str, ambig_slice):
    runs, cur_color, cur_text = [], None, ""
    for ch, is_red in zip(bits_str, ambig_slice):
        color = "red" if is_red else "black"
        if color != cur_color and cur_text:
            runs.append((cur_color, cur_text))
            cur_text = ""
        cur_color, cur_text = color, cur_text + ch
    if cur_text:
        runs.append((cur_color, cur_text))
    parts = [TextBlock(RED, t) if c == "red" else t for c, t in runs]
    return CellRichText(*parts) if parts else CellRichText("")


def append_row(ws, cells, bits_str, ambig_slice, fill):
    ws.append(cells)
    r = ws.max_row
    ws.cell(r, 5).value = bits_richtext(bits_str, ambig_slice)
    ws.cell(r, 5).font = Font(name="Consolas")
    ws.cell(r, 7).alignment = Alignment(wrap_text=True, vertical="top")
    for c in range(1, len(cells) + 1):
        ws.cell(r, c).fill = fill
    return r


def segment_fill_for_range(seg_label_slice):
    labels = sorted(set(lbl for lbl in seg_label_slice if lbl is not None))
    if len(labels) == 1:
        return SEGMENT_FILLS.get(labels[0], SEGMENT_FILLS["mixed"]), labels[0]
    return SEGMENT_FILLS["mixed"], "+".join(labels)


def write_workbook(out_path, frame_no, start, z, decode, fields, enums_ctx, audio_path):
    decided_bit = decode["decided_bit"]
    ambiguous = decode["ambiguous"]
    seg_label = decode["seg_label"]

    info_bits = decided_bit[INFO_BYTE0 * 8: (INFO_BYTE0 + INFO_NBYTES) * 8]
    info_ambig = ambiguous[INFO_BYTE0 * 8: (INFO_BYTE0 + INFO_NBYTES) * 8]
    info_seg = seg_label[INFO_BYTE0 * 8: (INFO_BYTE0 + INFO_NBYTES) * 8]
    n_avail = info_bits.size

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Frame{frame_no}_RawNoDestuff"

    align = decode["align"]

    headers = ["Subsystem", "ItemName", "OffsetBit", "BitLen", "Bits", "ReadableValue", "LookupTable"]
    legend = ("NO DESTUFFING APPLIED -- every row below is raw threshold=0/theta* symbol decisions, "
              "with stuffed bits still in the stream, so byte boundaries only match the true payload "
              "until the first real stuffed bit. Red text = data/header/FCS: near decision line "
              "(|y-theta*|<=" + f"{decode['margin']:.3f}" + ") | zero-run: decided-as-1 (should be 0x00). "
              "Row shading = segment (same meaning as the destuffed sheet, just naively byte-mapped): "
              "blue=header peach=data1 pink=data2 gray=zero-run1/2 purple=FCS green=trailing flags "
              "yellow=straddles a segment boundary.")
    ws.append([legend])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws.cell(1, 1).font = Font(italic=True, size=9)

    if align["offset_bits"] is None:
        banner = "Closing flag not found near the expected frame end in the raw stream."
        banner_font = Font(bold=True, size=9, color="FFCC0000")
    elif align["aligned"]:
        banner = ("RAW OFFSET = 0 bits: the closing flag lands exactly at raw bit "
                  f"{FRAME_END_BIT} even with no destuffing -- this frame has no real HDLC "
                  "stuffing in it at all.")
        banner_font = Font(bold=True, size=9, color="FF006600")
    else:
        banner = (f"RAW OFFSET = {align['offset_bits']:+d} bits: with NO destuffing applied, the "
                  f"closing flag 0x7E is found at raw bit {align['flag_run_start_bit']} instead of "
                  f"{FRAME_END_BIT} -- this is the total number of bits actually stuffed onto the "
                  f"wire before the closing flag, measured directly from the signal (compare against "
                  f"the destuffed sheet's align_slip_bits / align_stuffed_removed for the same frame).")
        banner_font = Font(bold=True, size=9, color="FFCC6600")
    ws.append([banner])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    ws.cell(2, 1).font = banner_font

    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(3, c)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    ws.freeze_panes = "A4"

    # ---- Header block (Dest/Src address, Control, PID) ----
    dest_bits = decided_bit[0:56]
    src_bits = decided_bit[56:112]
    ctrl_bits = decided_bit[112:120]
    pid_bits = decided_bit[120:128]

    def ascii_shift1(bits56):
        by = bits_to_bytes(bits56)
        callsign = "".join(chr(b >> 1) if 0x20 <= (b >> 1) < 0x7F else "?" for b in by[:6])
        return f"'{callsign}' (SSID byte=0x{by[6]:02X})"

    append_row(ws, ["HEADER", "Dest Address", 0, 56, None,
                    ascii_shift1(dest_bits), "expect 'BN0CU ' (>>1 of each byte, ground-test reference)"],
               "".join(str(int(b)) for b in dest_bits), ambiguous[0:56], SEGMENT_FILLS["header"])
    append_row(ws, ["HEADER", "Src Address", 56, 56, None,
                    ascii_shift1(src_bits), "expect 'BN0SCX' (>>1 of each byte, ground-test reference)"],
               "".join(str(int(b)) for b in src_bits), ambiguous[56:112], SEGMENT_FILLS["header"])
    ctrl_byte = bits_to_bytes(ctrl_bits)[0]
    append_row(ws, ["HEADER", "Control", 112, 8, None, f"0x{ctrl_byte:02X}", "expect 0x03 (UI frame)"],
               "".join(str(int(b)) for b in ctrl_bits), ambiguous[112:120], SEGMENT_FILLS["header"])
    pid_byte = bits_to_bytes(pid_bits)[0]
    append_row(ws, ["HEADER", "PID", 120, 8, None, f"0x{pid_byte:02X}", "expect 0xF0 (no layer-3)"],
               "".join(str(int(b)) for b in pid_bits), ambiguous[120:128], SEGMENT_FILLS["header"])

    # ---- Info field, one row per beacon field (OffsetBit order) ----
    n_missing = 0
    for f in fields:
        b0, b1 = f["offsetbit"], f["offsetbit"] + f["bitlen"]
        fill, seg_name = segment_fill_for_range(info_seg[b0:min(b1, n_avail)]) if b0 < n_avail \
            else (SEGMENT_FILLS["mixed"], "?")
        if b1 > n_avail:
            n_missing += 1
            bits_str = "?" * f["bitlen"]
            ambig_slice = np.zeros(f["bitlen"], dtype=bool)
            readable = "(no symbols recovered this far)"
            lookup_text = "-"
        else:
            bit_slice = info_bits[b0:b1]
            ambig_slice = info_ambig[b0:b1]
            bits_str = "".join(str(int(b)) for b in bit_slice)
            value, raw_unsigned = decode_raw_value(bits_str, f["dtype"], f["bitlen"], f["endian"])
            parsed_enum, transform, lookup_text = resolve_lookup(
                f["item_name"], f["longdesc"], f["bitlen"], enums_ctx)
            readable = format_readable(value, raw_unsigned, parsed_enum, transform)

        item_display = "(reserved)" if f["item_name"] == "NaN" else f["item_name"]
        append_row(ws, [f["subsystem"], item_display, f["offsetbit"], f["bitlen"], None,
                        readable, lookup_text], bits_str, ambig_slice, fill)

    # ---- FCS (2-byte CRC-16/X.25, at the naive un-shifted payload position) ----
    fcs_bits = decided_bit[272 * 8:274 * 8]
    fcs_ambig = ambiguous[272 * 8:274 * 8]
    fcs_bytes = bits_to_bytes(fcs_bits)
    transmitted_crc = int.from_bytes(fcs_bytes, "little")
    expected_crc = crc16_x25(bits_to_bytes(decided_bit[:272 * 8]))
    crc_note = (f"computed CRC16/X25(payload)=0x{expected_crc:04X} vs transmitted=0x{transmitted_crc:04X} "
                f"-> {'MATCH' if expected_crc == transmitted_crc else 'MISMATCH'}  "
                f"(expected to mismatch here whenever RAW OFFSET != 0, since real stuffed bits are still "
                f"in the payload)")
    append_row(ws, ["FCS", "FCS (CRC-16/X.25, little-endian)", 272 * 8, 16, None,
                    f"0x{transmitted_crc:04X}", crc_note],
               "".join(str(int(b)) for b in fcs_bits), fcs_ambig, SEGMENT_FILLS["FCS"])

    # ---- Trailing flags (known content: 0x7E repeated) ----
    expected_trail = trailing_expected_pattern(align)
    trail_thr0 = decode["bits_thr0"]
    for k in range(TRAIL_BYTES):
        b0, b1 = FRAME_END_BIT + k * 8, FRAME_END_BIT + k * 8 + 8
        if b1 > decided_bit.size:
            break
        bits_str = "".join(str(int(b)) for b in decided_bit[b0:b1])
        thr0_str = "".join(str(int(b)) for b in trail_thr0[b0:b1])
        ambig = ambiguous[b0:b1]
        e_theta = sum(1 for a, e in zip(bits_str, expected_trail) if a != e)
        e_thr0 = sum(1 for a, e in zip(thr0_str, expected_trail) if a != e)
        byte_val = bits_to_bytes(decided_bit[b0:b1])[0]
        note = (f"expect {expected_trail} (0x7E at the detected raw-offset phase) -> "
                f"{'OK' if e_theta == 0 else f'{e_theta}/8 differ'};  "
                f"same bits at plain threshold=0: {thr0_str} -> "
                f"{'OK' if e_thr0 == 0 else f'{e_thr0}/8 differ'}")
        append_row(ws, ["TRAILING", f"trailing flag byte {k}", b0, 8, None,
                        f"0x{byte_val:02X}", note],
                   bits_str, ambig, SEGMENT_FILLS["trailing-flag"])

    # ---- Summary ----
    seg_counts = decode["seg_counts"]
    total_ambig = sum(a for a, _ in seg_counts.values())
    total_bits = sum(n for _, n in seg_counts.values())
    breakdown = "  ".join(f"{lbl}={a}/{n}" for lbl, (a, n) in seg_counts.items())
    ws.append([])
    r = ws.max_row + 1
    ws.cell(r, 1).value = f"TOTAL possibly-wrong bits (red): {total_ambig} / {total_bits}"
    ws.cell(r, 1).font = Font(bold=True)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    ws.cell(r + 1, 1).value = f"Breakdown: {breakdown}"
    ws.cell(r + 1, 1).font = Font(italic=True, size=9)
    ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=7)
    ws.cell(r + 2, 1).value = banner
    ws.cell(r + 2, 1).font = Font(italic=True, size=9,
                                   color="FF006600" if align["aligned"] else "FFCC6600")
    ws.merge_cells(start_row=r + 2, start_column=1, end_row=r + 2, end_column=7)

    widths = {"A": 12, "B": 30, "C": 10, "D": 8, "E": 36, "F": 26, "G": 60}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    ws2 = wb.create_sheet("Frame_Info")
    info_rows = [
        ("frame_index", frame_no),
        ("start_sample", start),
        ("z_score", round(z, 3)),
        ("header_err_144", decode["hdr_err"]),
        ("offset_threshold", round(decode["offset_thr"], 5)),
        ("A_median_abs_y", round(decode["A"], 5)),
        ("margin_fraction", MARGIN_FRAC),
        ("margin_abs", round(decode["margin"], 5)),
        ("crc_pass_threshold0_baseline", decode["crc_baseline"]),
        ("crc_pass_mixed_threshold", decode["crc_mixed"]),
        ("z_threshold_used", Z_THRESHOLD),
        ("destuff_applied", False),
        ("raw_flag_run_start_bit", align["flag_run_start_bit"]),
        ("raw_expected_flag_bit", FRAME_END_BIT),
        ("raw_offset_bits", align["offset_bits"]),
        ("raw_aligned", align["aligned"]),
        ("info_bits_recovered", int(n_avail)),
        ("info_bits_expected", INFO_NBYTES * 8),
        ("fields_with_missing_bits", n_missing),
        ("total_ambiguous_bits", total_ambig),
        ("total_bits_all_segments", total_bits),
        ("audio_file", audio_path),
        ("xlsx_source", os.path.relpath(XLSX_PATH, HERE)),
        ("enums_source", os.path.relpath(ENUMS_PATH, HERE)),
    ]
    ws2.append(["key", "value"])
    for k, v in info_rows:
        ws2.append([k, v])
    ws2.append([])
    ws2.append(["segment", "ambiguous_bits", "total_bits"])
    for lbl, (a, n) in seg_counts.items():
        ws2.append([lbl, a, n])
    ws2.column_dimensions["A"].width = 28
    ws2.column_dimensions["B"].width = 50

    os.makedirs(OUT_DIR, exist_ok=True)
    wb.save(out_path)


def parse_args():
    parser = argparse.ArgumentParser(
        description="Same decode as 04_beacon_field_decode.py, but with destuffing skipped entirely.")
    parser.add_argument("audio", nargs="?", default=AUDIO,
                         help=f"path to the .ogg recording (default: {os.path.relpath(AUDIO, HERE)})")
    parser.add_argument("--frame", type=int, default=None,
                         help="1-based frame index to decode (default: decode every detected frame)")
    parser.add_argument("--z-threshold", type=float, default=Z_THRESHOLD,
                         help=f"frame-detection z-score threshold (default: {Z_THRESHOLD})")
    return parser.parse_args()


def main():
    global Z_THRESHOLD
    args = parse_args()
    Z_THRESHOLD = args.z_threshold

    with open(ENUMS_PATH, "r", encoding="utf-8") as fh:
        enums_ctx = json.load(fh)
    fields = load_beacon_fields(XLSX_PATH)
    print(f"Loaded {len(fields)} beacon fields from {os.path.basename(XLSX_PATH)}")
    print("NO DESTUFFING APPLIED -- this is the raw threshold=0/theta* view, stuffed bits still in the stream.")

    y, fs = audio_io.read_audio(args.audio, expected_fs=FS)
    bl = baseline.restore_baseline(y, num_iters=7, W=1000)
    yc = bl["y_comp_final"]

    template = build_header_template()
    starts, zs = detect_frame_starts(yc, template)
    print(f"Detected {len(starts)} frame(s) in {os.path.basename(args.audio)}: " +
          ", ".join(f"frame#{i}@{s} (z={z:.2f})" for i, (s, z) in enumerate(zip(starts, zs), start=1)))

    if args.frame is not None:
        if not (1 <= args.frame <= len(starts)):
            raise SystemExit(f"--frame {args.frame} out of range: only {len(starts)} frame(s) detected")
        frame_indices = [args.frame]
    else:
        frame_indices = list(range(1, len(starts) + 1))

    stem = os.path.splitext(os.path.basename(args.audio))[0]

    for frame_no in frame_indices:
        start, z = starts[frame_no - 1], zs[frame_no - 1]
        print(f"\n{'=' * 78}\nFrame#{frame_no} at sample={start} (z-score={z:.2f})\n{'=' * 78}")

        decode = decode_frame(y, yc, start)
        print(f"offset threshold theta*={decode['offset_thr']:+.4f}  header err={decode['hdr_err']}/144  "
              f"A={decode['A']:.4f}  margin=+/-{decode['margin']:.4f}")
        print(f"CRC (threshold=0 baseline): {'PASS' if decode['crc_baseline'] else 'fail'}   "
              f"CRC (mixed: theta* in data segs): {'PASS' if decode['crc_mixed'] else 'fail'}")

        seg_counts = decode["seg_counts"]
        total_ambig = sum(a for a, _ in seg_counts.values())
        total_bits = sum(n for _, n in seg_counts.values())
        print("Possibly-wrong (red) bits by segment: " +
              "  ".join(f"{lbl}={a}/{n}" for lbl, (a, n) in seg_counts.items()))
        print(f"TOTAL possibly-wrong bits: {total_ambig}/{total_bits}")

        a = decode["align"]
        if a["offset_bits"] is None:
            print("Raw offset: closing flag not found near the expected frame end.")
        elif a["aligned"]:
            print(f"Raw offset: 0 bits -- closing flag lands exactly at bit {FRAME_END_BIT} "
                  f"with NO destuffing applied (no real stuffing in this frame at all).")
        else:
            print(f"Raw offset: {a['offset_bits']:+d} bits -- closing flag found at raw bit "
                  f"{a['flag_run_start_bit']} vs expected {FRAME_END_BIT}, with NO destuffing applied. "
                  f"This is the total bits actually stuffed onto the wire, measured directly.")

        out_path = os.path.join(OUT_DIR, f"{stem}_frame{frame_no}_raw_no_destuff.xlsx")
        write_workbook(out_path, frame_no, start, z, decode, fields, enums_ctx,
                       os.path.relpath(args.audio, HERE))
        print(f"Workbook saved: {out_path}")


if __name__ == "__main__":
    main()
